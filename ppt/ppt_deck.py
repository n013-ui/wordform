# -*- coding: utf-8 -*-
"""簡報組：把 deck（Python dict）轉成「簡報內容與講稿.docx」與「簡報骨架.pptx」（講稿在備忘稿），並檢查規則。
用法（在 ChatGPT 的 Python 中）：
    exec(open('/mnt/data/ppt_deck.py', encoding='utf-8').read())
    print(make_deck(deck, '/mnt/data'))
需要 python-docx 與 python-pptx（ChatGPT 的 Python 環境已內建）。
"""
import re, os

FONT = '微軟正黑體'
NAMES = ['封面', '目錄', '前言', '內容頁', '內容頁', '內容頁', '內容頁', '內容頁', '心得結語', '參考網站']
TYPES = ['cover', 'toc', 'intro', 'content', 'content', 'content', 'content', 'content', 'conclusion', 'refs']
CATS = [
    ('數據類', ['bar chart', 'line chart', 'pie chart', 'stacked bar', 'gauge', 'barbell', 'grouped bar', 'donut', 'area chart']),
    ('流程類', ['flowchart', 'step-by-step', 'process']),
    ('時間類', ['timeline']),
    ('循環類', ['cycle']),
    ('比較類', ['side-by-side', 'comparison', 'pros and cons', 'before and after', "dos and don", 'venn']),
    ('階層與分類類', ['pyramid', 'hierarchy', 'pillars', 'three-column', 'column categories', 'tree']),
    ('關聯類', ['mind map', 'cause and effect', 'matrix', 'quadrant', 'funnel', 'iceberg']),
]
SHAPES = [
    ('欄位並排', ['side-by-side', 'three-column', 'column categories', 'pillars', 'pros and cons', 'comparison table']),
    ('由上往下分層', ['pyramid', 'hierarchy', 'funnel']),
    ('由左往右的線', ['flowchart', 'timeline', 'step-by-step']),
    ('長條', ['bar chart', 'stacked bar', 'grouped bar']),
]


def _norm(s):
    return re.sub(r'^(horizontal|vertical)\s+', '', str(s or '').lower()).strip()


def _pick(table, en):
    e = _norm(en)
    for name, keys in table:
        if any(k in e for k in keys):
            return name
    return None


def _hl_runs(text, hs):
    out, rest = [], str(text or '')
    hs = sorted([h for h in (hs or []) if h], key=len, reverse=True)
    while rest:
        best, word = -1, ''
        for h in hs:
            k = rest.find(h)
            if k >= 0 and (best < 0 or k < best):
                best, word = k, h
        if best < 0:
            out.append((rest, False)); break
        if best > 0:
            out.append((rest[:best], False))
        out.append((word, True)); rest = rest[best + len(word):]
    return out


def check(deck):
    """回傳 (錯誤清單, 建議清單, 圖表一覽, 講稿統計)"""
    fails, warns, overview = [], [], []
    S = deck.get('slides', [])
    if len(S) != 10:
        return [f'投影片要剛好 10 頁（目前 {len(S)} 頁）'], [], [], {}
    for i, s in enumerate(S):
        if s.get('type') != TYPES[i]:
            fails.append(f"第 {i + 1} 頁的 type 應該是 {TYPES[i]}（目前 {s.get('type')}）")
    if fails:
        return fails, [], [], {}
    cover, intro, concl, refs = S[0], S[2], S[8], S[9]
    if len(cover.get('title', '')) > 14: warns.append('封面標題建議 12 字以內')
    if len(cover.get('subtitle', '')) > 22: warns.append('副標題建議 20 字以內')
    for x in (intro, concl):
        nm = '前言' if x['type'] == 'intro' else '結語'
        if x.get('paragraph'):
            n = len(x['paragraph'])
            if n < 70 or n > 170: warns.append(f'{nm}短文 {n} 字，建議 80～150 字')
            hs = x.get('highlights', [])
            if not 3 <= len(hs) <= 6: warns.append(f'{nm}的重點詞 {len(hs)} 個，建議 3～5 個')
            for h in hs:
                if h not in x['paragraph']: fails.append(f'{nm}的重點詞「{h}」沒有出現在短文裡')
    n_ref = len(refs.get('refs', []))
    if not 3 <= n_ref <= 6: fails.append(f'參考網站要 3～6 筆（目前 {n_ref} 筆）')
    used_cat, used_shape = {}, {}
    for k, s in enumerate(S[3:8]):
        no, d = k + 4, s.get('diagram', {})
        en = d.get('en')
        if not en:
            fails.append(f'第 {no} 頁沒有圖形（diagram.en）'); continue
        cat, shape = _pick(CATS, en), _pick(SHAPES, en)
        if not cat: warns.append(f'第 {no} 頁的圖形「{en}」不在對照表中，請確認名稱')
        if cat:
            if cat in used_cat: fails.append(f'第 {used_cat[cat]} 頁和第 {no} 頁都是{cat}，五張圖要來自五個不同類別')
            else: used_cat[cat] = no
        if shape:
            if shape in used_shape: fails.append(f'第 {used_shape[shape]} 頁和第 {no} 頁的圖形狀相似（都是「{shape}」），請換一種')
            else: used_shape[shape] = no
        lines = [x for x in d.get('text', []) if str(x).strip()]
        if not lines or not re.match(r'^[A-Za-z].*[:：]', lines[0]):
            fails.append(f'第 {no} 頁 Napkin 文字的第一行要是「English name: 圖的標題」')
        if not 3 <= len(lines[1:]) <= 7: warns.append(f'第 {no} 頁 Napkin 條列 {len(lines[1:])} 項，建議 3～6 項')
        if len(s.get('title', '')) > 14: warns.append(f'第 {no} 頁標題建議 12 字以內')
        overview.append(f"第{no}頁 {en}（{cat or '?'}{'／' + shape if shape else ''}）")
    for k, s in enumerate(S[3:8]):
        no, f = k + 4, s.get('diagram', {}).get('fallback')
        if not f: continue
        fc, fs = _pick(CATS, f), _pick(SHAPES, f)
        if fc and fc in used_cat and used_cat[fc] != no: warns.append(f'第 {no} 頁的備用圖形 {f} 和第 {used_cat[fc]} 頁同屬{fc}，建議換一個')
        elif fs and fs in used_shape and used_shape[fs] != no: warns.append(f'第 {no} 頁的備用圖形 {f} 和第 {used_shape[fs]} 頁形狀相似，建議換一個')
    minutes = float(deck.get('minutes') or 7)
    script = deck.get('script', [])
    if len(script) != 10: fails.append(f'講稿要對應 10 頁（目前 {len(script)} 頁）')
    total, secs, by = 0, 0, {'A': 0, 'B': 0}
    for i, p in enumerate(script):
        ls = p.get('lines', [])
        secs += float(p.get('seconds') or 0)
        for who, t in ls:
            total += len(t); by[who] = by.get(who, 0) + len(t)
        if 3 <= i <= 7:
            if len(ls) < 4: fails.append(f'第 {i + 1} 頁講稿只有 {len(ls)} 句，內容頁至少 4 句來回')
            if len({w for w, _ in ls}) < 2: fails.append(f'第 {i + 1} 頁講稿只有一個人說話')
    target = minutes * 250
    if total < target * 0.85: fails.append(f'講稿共 {total} 字，{minutes:g} 分鐘約需 {target:.0f} 字，請補充解釋與例子')
    elif total > target * 1.15: fails.append(f'講稿共 {total} 字，超過 {minutes:g} 分鐘（約 {target:.0f} 字），請精簡')
    if secs and abs(secs - minutes * 60) > minutes * 6: warns.append(f'各頁秒數加總 {secs:.0f} 秒，和 {minutes:g} 分鐘差太多')
    share = by.get('A', 0) / total if total else 0.5
    if not 0.38 <= share <= 0.62: fails.append(f'A 和 B 的發言量不平均（A 約 {share * 100:.0f}%）')
    return fails, warns, overview, {'total': total, 'target': target, 'share': share, 'secs': secs, 'minutes': minutes}


def _docx(deck, path, overview):
    from docx import Document
    from docx.shared import Pt, RGBColor, Cm
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.2); sec.top_margin = sec.bottom_margin = Cm(2)
    st = doc.styles['Normal']; st.font.name = FONT; st.font.size = Pt(12)
    st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), FONT)

    def para(runs, indent=0, before=0, after=4):
        p = doc.add_paragraph()
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(before), Pt(after)
        if indent: p.paragraph_format.left_indent = Pt(indent)
        for t, o in runs:
            r = p.add_run(t); r.font.size = Pt(o.get('size', 12)); r.bold = o.get('bold', False)
            r.font.name = FONT; r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), FONT)
            if o.get('color'): r.font.color.rgb = RGBColor.from_string(o['color'])
        return p

    def head(t, lvl):
        p = para([(t, {'bold': True, 'size': 16 if lvl == 1 else 13, 'color': '1F4E79' if lvl == 1 else '1F2933'})], before=12 if lvl == 1 else 8)
        p.paragraph_format.keep_with_next = True

    def kv(k, v):
        para([(k + '：', {'bold': True}), (str(v), {})])

    def box(lines):
        t = doc.add_table(rows=1, cols=1); c = t.rows[0].cells[0]
        sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), 'F4F6F8')
        c._tc.get_or_add_tcPr().append(sh)
        c.paragraphs[0].text = ''
        for k, l in enumerate(lines):
            p = c.paragraphs[0] if k == 0 else c.add_paragraph()
            r = p.add_run(str(l)); r.font.size = Pt(11); r.font.name = FONT
            r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), FONT)

    S, script, minutes = deck['slides'], deck.get('script', []), deck.get('minutes', 7)
    para([(f"簡報組｜{deck.get('topic', '')}", {'bold': True, 'size': 18})])
    para([(f'簡報時間 {minutes} 分鐘｜兩人對話（A、B）｜第 4～8 頁的圖用 Napkin AI 製作', {'size': 10, 'color': '5D6878'})])
    if deck.get('overview'):
        head('解題重點', 1)
        for t in deck['overview']: para([('・' + t, {})])
    head('一、簡報內容', 1)
    for i, s in enumerate(S):
        head(f"第 {i + 1} 頁｜{NAMES[i]}{'：' + s['title'] if s.get('title') and 3 <= i <= 7 else ''}", 2)
        tp = s['type']
        if tp == 'cover': kv('標題', s.get('title', '')); kv('副標題', s.get('subtitle', ''))
        if tp == 'toc':
            for k, t in enumerate(s.get('items', [])): para([(f'{k + 1}. {t}', {})], indent=18)
        if tp in ('intro', 'conclusion'):
            if s.get('paragraph'):
                para([(t, {'bold': True, 'color': 'C00000'} if h else {}) for t, h in _hl_runs(s['paragraph'], s.get('highlights'))], indent=12)
                kv('重點詞（投影片上用紅色粗體）', '、'.join(s.get('highlights', [])))
            for t in s.get('lines', []): para([('・' + t, {})], indent=12)
            if s.get('visual'): kv('建議畫面', s['visual'])
        if tp == 'content':
            d = s.get('diagram', {})
            kv('一句話重點', s.get('key', ''))
            kv('圖形類型', f"{d.get('zh', '')}（{d.get('en', '')}，{_pick(CATS, d.get('en')) or ''}）")
            if d.get('why'): kv('為什麼用這種圖', d['why'])
            kv('Custom Generation 輸入框請打', d.get('en', ''))
            para([('貼進 Napkin 的文字：', {'bold': True})])
            box(d.get('text', []))
            if d.get('fallback'): kv('備用圖形', d['fallback'])
            if d.get('source'): kv('資料來源', '第 10 頁 ' + str(d['source']))
        if tp == 'refs':
            for k, x in enumerate(s.get('refs', [])): para([(f"{k + 1}. {x['name']}：{x['url']}", {'size': 11})], indent=18)
    head('圖表多樣性檢查', 2)
    for t in overview: para([('・' + t, {})])
    doc.add_page_break()
    head('二、講稿（兩人對話）', 1)
    for i, p in enumerate(script):
        head(f"第 {i + 1} 頁｜{NAMES[i]}（約 {p.get('seconds', '?')} 秒）", 2)
        for who, t in p.get('lines', []):
            para([(who + '：', {'bold': True, 'color': '1F5F99' if who == 'A' else 'B4541A'}), (t, {})], indent=12)
    if deck.get('tips'):
        head('上台小提醒', 2)
        for t in deck['tips']: para([('・' + t, {})])
    doc.save(path)


def _pptx(deck, path):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.enum.dml import MSO_LINE_DASH_STYLE
    INK, BRAND, MUTED, SOFT, LINE = '1F2933', '1F5F99', '5D6878', 'EEF3F9', 'C9D6E6'
    prs = Presentation(); prs.slide_width, prs.slide_height = Inches(10), Inches(5.625)
    blank = prs.slide_layouts[6]

    def rect(sl, x, y, w, h, fill, line=None, dash=False):
        s = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        s.fill.solid(); s.fill.fore_color.rgb = RGBColor.from_string(fill)
        s.line.color.rgb = RGBColor.from_string(line or fill)
        if dash: s.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        s.shadow.inherit = False
        return s

    def text(sl, x, y, w, h, runs_list, size=18, color=INK, bold=False, align=None, anchor=MSO_ANCHOR.TOP):
        tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame; tf.word_wrap = True
        tf.vertical_anchor = anchor
        for k, runs in enumerate(runs_list):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            if align is not None: p.alignment = align
            p.space_after = Pt(8)
            for t, o in runs:
                r = p.add_run(); r.text = t; f = r.font
                f.name = FONT; f.size = Pt(o.get('size', size)); f.bold = o.get('bold', bold)
                f.color.rgb = RGBColor.from_string(o.get('color', color))
        return tb

    def header(sl, t, no):
        rect(sl, 0, 0, 10, 0.12, BRAND)
        text(sl, 0.5, 0.3, 8.6, 0.7, [[(t, {})]], size=28, bold=True)
        text(sl, 9.2, 5.15, 0.6, 0.3, [[(str(no), {})]], size=11, color=MUTED, align=PP_ALIGN.RIGHT)

    S, script = deck['slides'], deck.get('script', [])
    for i, s in enumerate(S):
        sl = prs.slides.add_slide(blank); tp = s['type']
        if tp == 'cover':
            rect(sl, 0, 0, 10, 5.625, SOFT); rect(sl, 0.6, 2.55, 1.2, 0.08, BRAND)
            text(sl, 0.6, 1.2, 8.8, 1.2, [[(s.get('title', ''), {})]], size=40, bold=True)
            text(sl, 0.6, 2.8, 8.8, 0.8, [[(s.get('subtitle', ''), {})]], size=22, color=MUTED)
            text(sl, 0.6, 4.5, 8.8, 0.5, [[('報告人：○○○、○○○', {})]], size=16, color=MUTED)
        elif tp == 'toc':
            header(sl, '目錄', i + 1)
            text(sl, 0.9, 1.2, 8.2, 4, [[(f'{k + 1:02d}　{t}', {})] for k, t in enumerate(s.get('items', []))], size=22)
        elif tp in ('intro', 'conclusion'):
            header(sl, '前言' if tp == 'intro' else '心得與結語', i + 1)
            w = 5.2 if s.get('visual') else 8.8
            if s.get('paragraph'):
                text(sl, 0.6, 1.25, w, 3.8, [[(t, {'bold': True, 'color': 'C00000'} if h else {}) for t, h in _hl_runs(s['paragraph'], s.get('highlights'))]], size=18)
            else:
                text(sl, 0.6, 1.25, w, 3.8, [[('・' + t, {})] for t in s.get('lines', [])], size=22)
            if s.get('visual'):
                rect(sl, 6.1, 1.3, 3.4, 3.6, SOFT, LINE, dash=True)
                text(sl, 6.25, 1.4, 3.1, 3.4, [[('建議畫面：' + s['visual'], {})]], size=13, color=MUTED, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        elif tp == 'content':
            d = s.get('diagram', {})
            header(sl, s.get('title', ''), i + 1)
            text(sl, 0.5, 0.95, 9, 0.5, [[(s.get('key', ''), {})]], size=16, color=BRAND, bold=True)
            rect(sl, 0.5, 1.55, 9, 3.55, 'FAFBFC', LINE, dash=True)
            first = (d.get('text') or [''])[0]
            text(sl, 0.7, 1.7, 8.6, 3.25, [[(f"【Napkin 圖：{d.get('en', '')}】", {'bold': True})], [(first, {})], [('（在 Napkin 生成後，刪除此框並貼上圖片）', {})]],
                 size=14, color=MUTED, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        elif tp == 'refs':
            header(sl, '參考資料', i + 1)
            text(sl, 0.6, 1.2, 8.8, 4, [[(f"{k + 1}. {x['name']}：{x['url']}", {})] for k, x in enumerate(s.get('refs', []))], size=12)
        if i < len(script):
            sl.notes_slide.notes_text_frame.text = '\n'.join(f'{w}：{t}' for w, t in script[i].get('lines', []))
    prs.save(path)


def make_deck(deck, out_dir='/mnt/data'):
    """主程式：deck＝簡報內容（dict），out_dir＝輸出資料夾。回傳檢查報告文字。"""
    fails, warns, overview, st = check(deck)
    if fails and not overview and not st:
        return '\n'.join('✗ ' + f for f in fails) + '\n檢查結果：需要修改（格式錯誤，未產生檔案）'
    os.makedirs(out_dir, exist_ok=True)
    base = re.sub(r'[\\/:*?"<>|？]', '', deck.get('topic', ''))[:16]
    base = base + '_' if base else ''
    dp = os.path.join(out_dir, base + '簡報內容與講稿.docx')
    pp = os.path.join(out_dir, base + '簡報骨架.pptx')
    _docx(deck, dp, overview)
    _pptx(deck, pp)
    out = [f"講稿：共 {st['total']} 字（{st['minutes']:g} 分鐘約需 {st['target']:.0f} 字），A {st['share'] * 100:.0f}%／B {100 - st['share'] * 100:.0f}%，各頁秒數合計 {st['secs']:.0f} 秒",
           '圖表：' + '、'.join(overview)]
    out += ['△ ' + w for w in warns] + ['✗ ' + f for f in fails]
    out += ['已產生 ' + dp, '已產生 ' + pp + '（講稿在每頁的備忘稿）', '檢查結果：' + ('需要修改' if fails else '通過')]
    return '\n'.join(out)


import sys as _sys
if __name__ == '__main__' and _sys.argv and _sys.argv[0].endswith('ppt_deck.py'):
    import json
    if len(_sys.argv) < 3:
        print('用法：python3 ppt_deck.py deck.json 輸出資料夾'); _sys.exit(2)
    rep = make_deck(json.load(open(_sys.argv[1], encoding='utf-8')), _sys.argv[2])
    print(rep); _sys.exit(0 if '檢查結果：通過' in rep else 1)
