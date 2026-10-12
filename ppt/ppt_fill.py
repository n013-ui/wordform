# -*- coding: utf-8 -*-
"""簡報組：把簡報內容（deck）填進學生自己的 PPT 範本，保留範本的字型、顏色、背景與版面。
用法（Python）：
    exec(open('ppt_fill.py', encoding='utf-8').read())
    print(inspect_template('範本.pptx'))                 # 先看範本每一頁被認成什麼
    print(fill_template('範本.pptx', deck, '輸出.pptx'))  # 自動填入，回傳報告
deck 的格式與 ppt_deck.py 相同（slides 10 頁＋script 10 段）。需要 python-pptx。

自動辨識規則：
  第一頁＝封面；有「目錄」＝目錄；標題含「前言」＝前言；含「心得／結語／結論」＝結語；
  含「資料來源／參考」＝參考資料；其餘有標題＋內文的頁面＝內容頁；看不出來的頁面不動，並列在報告中。
  內容頁不足 5 頁會複製最後一個內容頁補足；超過 5 頁會刪掉多的。講稿放進每頁的備忘稿。
"""
import copy, re
from lxml import etree

NS = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
      'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
      'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
A = '{%s}' % NS['a']
RED = 'C00000'
EMU_IN = 914400


# ---------------- 讀取範本 ----------------
def _text_shapes(shapes, out=None):
    """遞迴找出所有有文字框的圖形（含群組內）"""
    out = [] if out is None else out
    for sh in shapes:
        if sh.shape_type == 6:          # 群組
            _text_shapes(sh.shapes, out)
        elif sh.has_text_frame:
            out.append(sh)
    return out


def _font_size(sh):
    sizes = []
    for p in sh.text_frame.paragraphs:
        for r in p.runs:
            if r.font.size: sizes.append(r.font.size.pt)
    if not sizes:
        for e in sh._element.iter(A + 'defRPr', A + 'endParaRPr'):
            if e.get('sz'): sizes.append(int(e.get('sz')) / 100)
    return max(sizes) if sizes else 18


def _abs_pos(sh):
    """群組內圖形的座標換算成投影片座標（英吋）"""
    x, y, w, h = sh.left or 0, sh.top or 0, sh.width or 0, sh.height or 0
    el = sh._element.getparent()
    while el is not None and el.tag.endswith('}grpSp'):
        xfrm = el.find('.//' + A + 'xfrm')
        if xfrm is not None:
            off, ext = xfrm.find(A + 'off'), xfrm.find(A + 'ext')
            choff, chext = xfrm.find(A + 'chOff'), xfrm.find(A + 'chExt')
            if None not in (off, ext, choff, chext) and int(chext.get('cx')) and int(chext.get('cy')):
                sx = int(ext.get('cx')) / int(chext.get('cx')); sy = int(ext.get('cy')) / int(chext.get('cy'))
                x = int(off.get('x')) + (x - int(choff.get('x'))) * sx
                y = int(off.get('y')) + (y - int(choff.get('y'))) * sy
                w, h = w * sx, h * sy
        el = el.getparent()
    return x / EMU_IN, y / EMU_IN, w / EMU_IN, h / EMU_IN


def _analyze(prs):
    """回傳每頁的角色與重要圖形"""
    info = []
    for i, s in enumerate(prs.slides):
        shs = [sh for sh in _text_shapes(s.shapes)]
        texts = [(sh, sh.text_frame.text.strip()) for sh in shs]
        nonempty = [(sh, t) for sh, t in texts if t]
        title = None
        if nonempty:
            title = sorted(nonempty, key=lambda st: (-_font_size(st[0]), _abs_pos(st[0])[1]))[0][0]
        tt = title.text_frame.text.strip() if title is not None else ''
        alltext = ' '.join(t for _, t in nonempty)
        bodies = [sh for sh, t in nonempty if sh is not title]
        if '@' in alltext and not re.search(r'[一-鿿]', alltext):
            role = 'other'
        elif re.search(r'目\s*錄|CONTENTS?|大綱', tt, re.I):
            role = 'toc'
        elif i == 0:
            role = 'cover'
        elif '前言' in tt or '緣起' in tt or '動機' in tt:
            role = 'intro'
        elif re.search(r'心得|結語|結論|總結', tt):
            role = 'conclusion'
        elif re.search(r'資料來源|參考|來源|引用', tt):
            role = 'refs'
        elif title is not None and bodies:
            role = 'content'
        else:
            role = 'other'
        info.append({'index': i, 'role': role, 'title': title, 'bodies': bodies, 'shapes': shs, 'title_text': tt})
    return info


def inspect_template(path):
    from pptx import Presentation
    prs = Presentation(path)
    lines = [f'範本共 {len(prs.slides)} 頁，大小 {prs.slide_width / EMU_IN:.2f} x {prs.slide_height / EMU_IN:.2f} 吋']
    ROLE = {'cover': '封面', 'toc': '目錄', 'intro': '前言', 'content': '內容頁', 'conclusion': '心得結語', 'refs': '參考資料', 'other': '無法辨識（不更動）'}
    for d in _analyze(prs):
        boxes = '、'.join('「%s」' % sh.text_frame.text.strip().replace('\n', ' ')[:12] for sh in d['shapes'] if sh.text_frame.text.strip())
        lines.append(f"第 {d['index'] + 1} 頁：{ROLE[d['role']]}　文字框：{boxes[:120]}")
    return '\n'.join(lines)


# ---------------- 寫入文字（保留範本格式） ----------------
def _set_runs(sh, paras, size=None, autofit=False):
    """paras：[[(文字, {'bold':..., 'color':...}), ...], ...]。沿用原本第一段、第一個文字的格式。"""
    tf = sh.text_frame
    txBody = tf._txBody
    ps = txBody.findall(A + 'p')
    base_p = ps[0] if ps else etree.SubElement(txBody, A + 'p')
    base_ppr = base_p.find(A + 'pPr')
    base_r = base_p.find(A + 'r')
    base_rpr = base_r.find(A + 'rPr') if base_r is not None else None
    if base_rpr is None:
        e = base_p.find(A + 'endParaRPr')
        base_rpr = copy.deepcopy(e) if e is not None else etree.Element(A + 'rPr')
        base_rpr.tag = A + 'rPr'
    for p in ps:
        txBody.remove(p)
    for runs in paras:
        p = etree.SubElement(txBody, A + 'p')
        if base_ppr is not None:
            p.append(copy.deepcopy(base_ppr))
        for text, o in runs:
            r = etree.SubElement(p, A + 'r')
            rpr = copy.deepcopy(base_rpr); rpr.tag = A + 'rPr'
            for k in ('dirty', 'err'):
                if k in rpr.attrib: del rpr.attrib[k]
            if size: rpr.set('sz', str(int(size * 100)))
            if o.get('bold'): rpr.set('b', '1')
            if o.get('color'):
                for c in rpr.findall(A + 'solidFill'): rpr.remove(c)
                fill = etree.Element(A + 'solidFill'); clr = etree.SubElement(fill, A + 'srgbClr'); clr.set('val', o['color'])
                # solidFill 要放在 latin/ea 字型設定前面
                pos = 0
                for k, ch in enumerate(rpr):
                    if ch.tag in (A + 'ln',): pos = k + 1
                rpr.insert(pos, fill)
            r.append(rpr)
            t = etree.SubElement(r, A + 't'); t.text = text
    tf.word_wrap = True
    if autofit:
        bp = txBody.find(A + 'bodyPr')
        if bp is not None:
            for k in list(bp):
                if k.tag in (A + 'spAutoFit', A + 'normAutofit', A + 'noAutofit'): bp.remove(k)
            etree.SubElement(bp, A + 'normAutofit')


def _hl_runs(text, hs):
    out, rest = [], str(text or '')
    hs = sorted([h for h in (hs or []) if h], key=len, reverse=True)
    while rest:
        best, word = -1, ''
        for h in hs:
            k = rest.find(h)
            if k >= 0 and (best < 0 or k < best): best, word = k, h
        if best < 0: out.append((rest, {})); break
        if best > 0: out.append((rest[:best], {}))
        out.append((word, {'bold': True, 'color': RED})); rest = rest[best + len(word):]
    return out


def _fit_size(text, w_in, h_in, start, lowest=14, line=1.35):
    """在 w x h 吋的框內放得下 text 的最大字級（中文字寬＝字級）"""
    n = len(text)
    for sz in [s for s in (start, 28, 26, 24, 22, 20, 18, 16, 14) if s <= start and s >= lowest]:
        per_line = max(1, int(w_in * 72 / sz))
        lines = sum(max(1, -(-len(seg) // per_line)) for seg in text.split('\n'))
        if lines * sz * line / 72 <= h_in:
            return sz
    return lowest


def _text_w(text, size):
    return sum(size / 72 * (0.55 if ord(c) < 127 else 1.0) for c in text)


def _fit_line(sh, text, others, W, margin=0.5, min_size=14, allow_center=False):
    """單行文字：把框加寬到放得下（避開同一列右邊的其他文字框），還放不下才縮小字級。回傳字級。"""
    x, y, w, h = _abs_pos(sh)
    size = _font_size(sh)
    need = _text_w(text, size) * 1.08 + 0.4          # 加上文字框內距與字寬誤差
    centered = allow_center and abs((x + w / 2) - W / 2) < 1.2 and w < W * 0.6
    if centered:
        avail = W - 2 * margin
    else:
        right = W - margin
        for o in others:
            if o is sh:
                continue
            ox, oy, ow, oh = _abs_pos(o)
            if ox > x + 0.05 and oy < y + h and oy + oh > y and o.text_frame.text.strip():
                right = min(right, ox - 0.15)
        avail = max(right - x, w)
    if need > avail:
        size = max(min_size, size * avail / need)
        need = avail
    if need > w:
        scale = w / ((sh.width or 1) / EMU_IN) if sh.width else 1
        new_w = need / (scale or 1)
        if centered:
            sh.left = int((sh.left or 0) - (new_w - (sh.width or 0) / EMU_IN) * EMU_IN / 2)
        sh.width = int(new_w * EMU_IN)
    return round(size, 1)


def _place(sh, x, y, w, h):
    from pptx.util import Inches
    sh.left, sh.top, sh.width, sh.height = Inches(x), Inches(y), Inches(w), Inches(h)


# ---------------- 投影片複製、刪除、排序 ----------------
def _duplicate(prs, src_index, insert_at):
    src = prs.slides[src_index]
    new = prs.slides.add_slide(src.slide_layout)
    for shp in list(new.shapes):
        shp._element.getparent().remove(shp._element)
    rid_map = {}
    for rid, rel in src.part.rels.items():
        if rel.is_external or 'slideLayout' in rel.reltype or 'notesSlide' in rel.reltype:
            continue
        rid_map[rid] = new.part.relate_to(rel._target, rel.reltype)
    sp_tree = new.shapes._spTree
    for el in src.shapes._spTree:
        if el.tag.endswith('}nvGrpSpPr') or el.tag.endswith('}grpSpPr'):
            continue
        c = copy.deepcopy(el)
        for e in c.iter():
            for attr in ('{%s}embed' % NS['r'], '{%s}id' % NS['r'], '{%s}link' % NS['r']):
                if e.get(attr) in rid_map: e.set(attr, rid_map[e.get(attr)])
        sp_tree.append(c)
    bg = src._element.cSld.find('{%s}bg' % NS['p'])
    if bg is not None:
        new._element.cSld.insert(0, copy.deepcopy(bg))
    ids = prs.slides._sldIdLst
    el = list(ids)[-1]; ids.remove(el); ids.insert(insert_at, el)
    return new


def _delete(prs, index):
    ids = prs.slides._sldIdLst
    el = list(ids)[index]
    prs.part.drop_rel(el.get('{%s}id' % NS['r']))
    ids.remove(el)


# ---------------- 主程式 ----------------
def fill_template(template_path, deck, out_path):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.dml import MSO_LINE_DASH_STYLE
    from pptx.dml.color import RGBColor
    prs = Presentation(template_path)
    W, H = prs.slide_width / EMU_IN, prs.slide_height / EMU_IN
    S, script = deck['slides'], deck.get('script', [])
    contents = [s for s in S if s.get('type') == 'content']
    report = []

    info = _analyze(prs)
    roles = [d['role'] for d in info]
    # 內容頁數量調整成和 deck 一樣
    cidx = [i for i, r in enumerate(roles) if r == 'content']
    if not cidx:
        return '✗ 範本裡找不到內容頁（有標題與內文框的頁面），無法自動填入。'
    while len(cidx) < len(contents):
        last = cidx[-1]
        _duplicate(prs, last, last + 1)
        report.append(f'範本只有 {len(cidx)} 個內容頁，複製第 {last + 1} 頁補到 {len(cidx) + 1} 頁')
        info = _analyze(prs); roles = [d['role'] for d in info]
        cidx = [i for i, r in enumerate(roles) if r == 'content']
    while len(cidx) > len(contents):
        _delete(prs, cidx[-1])
        report.append(f'範本的內容頁比需要的多，刪除原本的第 {cidx[-1] + 1} 頁')
        info = _analyze(prs); roles = [d['role'] for d in info]
        cidx = [i for i, r in enumerate(roles) if r == 'content']

    by_type = {s['type']: s for s in S if s['type'] != 'content'}
    ci = 0
    margin = 0.6
    for d in info:
        s_obj = prs.slides[d['index']]
        role, title, bodies = d['role'], d['title'], d['bodies']
        k = None
        if role == 'cover':
            c = by_type.get('cover', {}); k = 0
            if title is not None:
                _set_runs(title, [[(c.get('title', ''), {})]], size=_fit_line(title, c.get('title', ''), d['shapes'], W, allow_center=True))
            others = sorted([b for b in bodies], key=lambda b: -_font_size(b))
            if others:
                _set_runs(others[0], [[(c.get('subtitle', ''), {})]], size=_fit_line(others[0], c.get('subtitle', ''), d['shapes'], W, allow_center=True))
            for b in others[1:]:
                if re.search(r'報告|姓名|組員|作者', b.text_frame.text): pass
        elif role == 'toc':
            k = 1
            mids = [c.get('title', '') for c in contents]      # 目錄的「標題」位置依序放五個內容頁標題
            slots = [b for b in bodies if re.fullmatch(r'\s*(標題|主題|內容|小標|項目|標題\s*\d*)\s*', b.text_frame.text)]
            nums = []
            for b in [x for x in d['shapes'] if re.fullmatch(r'\s*0?\d\s*', x.text_frame.text)]:
                x, y, _, _ = _abs_pos(b); nums.append((int(b.text_frame.text.strip()), x, y))

            def order_key(b):
                x, y, _, _ = _abs_pos(b)
                if nums:
                    n = min(nums, key=lambda t: (t[1] - x) ** 2 + (t[2] - y) ** 2)
                    return (n[0], y, x)
                return (0, y, x)
            slots.sort(key=order_key)
            sizes = [_fit_line(b, t, d['shapes'], W) for b, t in zip(slots, mids)]
            # 目錄各項用同一個字級，看起來比較整齊；但不小於 26 點（空間不夠的那一項會自動換行）
            common = max(26, min(sizes)) if sizes else None
            common = min(common, max(_font_size(b) for b in slots)) if slots else common
            for b, t in zip(slots, mids):
                _set_runs(b, [[(t, {})]], size=common)
            if len(slots) != len(mids):
                report.append(f'△ 目錄有 {len(slots)} 個標題位置，但內容頁有 {len(mids)} 頁，請自行調整目錄')
            # 目錄編號重新排成連續的 1、2、3…（依「前言 → 內容頁 → 心得 → 資料來源」的順序）
            num_shapes = [x for x in d['shapes'] if re.fullmatch(r'\s*0?\d\s*', x.text_frame.text)]
            labels = [x for x in d['shapes'] if x not in num_shapes and x is not title and x.text_frame.text.strip()]

            def rank(lb):
                t = lb.text_frame.text
                if '前言' in t: return (0, 0)
                if lb in slots: return (1, slots.index(lb))
                if re.search(r'心得|結語|結論', t): return (2, 0)
                if re.search(r'來源|參考', t): return (3, 0)
                return (4, 0)
            if num_shapes and labels:
                pairs = []
                for nsh in num_shapes:
                    nx, ny, _, _ = _abs_pos(nsh)
                    lb = min(labels, key=lambda l: ((_abs_pos(l)[0] - nx) ** 2 + (_abs_pos(l)[1] - ny) ** 2))
                    pairs.append((rank(lb), nsh))
                pairs.sort(key=lambda t: t[0])
                old_nums = [int(n.text_frame.text.strip()) for _, n in pairs]
                for k2, (_, nsh) in enumerate(pairs, 1):
                    _set_runs(nsh, [[(str(k2), {})]])
                if old_nums != list(range(1, len(pairs) + 1)):
                    report.append('目錄編號已重新排成 1～%d（範本原本的編號有重複或不連續）' % len(pairs))
        elif role in ('intro', 'conclusion'):
            c = by_type.get(role, {}); k = 2 if role == 'intro' else 8
            body = max(bodies, key=lambda b: (b.width or 0) * (b.height or 0)) if bodies else None
            if body is not None:
                tx, ty, _, _ = _abs_pos(title) if title is not None else (margin, 0.4, 0, 0)
                top = ty + (title.height / EMU_IN if title is not None else 0.9) + 0.4
                x0 = max(margin, min(_abs_pos(body)[0], 2.5))
                w, h = W - x0 - margin, H - top - margin
                if c.get('paragraph'):
                    text = c['paragraph']; runs = [_hl_runs(text, c.get('highlights'))]
                else:
                    text = '\n'.join(c.get('lines', [])); runs = [[('・' + t, {})] for t in c.get('lines', [])]
                sz = _fit_size(text, w, h, min(_font_size(body), 28))
                _place(body, x0, top, w, h)
                _set_runs(body, runs, size=sz)
        elif role == 'content':
            c = contents[ci]; k = 3 + ci; ci += 1
            dg = c.get('diagram', {})
            if title is not None:
                _set_runs(title, [[(c.get('title', ''), {})]], size=_fit_line(title, c.get('title', ''), d['shapes'], W))
            body = max(bodies, key=lambda b: (b.width or 0) * (b.height or 0)) if bodies else None
            tx, ty, _, th = _abs_pos(title) if title is not None else (margin, 0.4, 0, 0.9)
            top = ty + th + 0.25
            if body is not None:
                x0 = max(margin, min(_abs_pos(body)[0], 2.5))
                _place(body, x0, top, W - x0 - margin, 0.7)
                _set_runs(body, [[(c.get('key', ''), {'bold': True})]], size=min(_font_size(body), 24))
            box_top = top + 0.85
            rect = s_obj.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(margin + 0.3), Inches(box_top), Inches(W - 2 * margin - 0.6), Inches(max(1.5, H - box_top - margin)))
            rect.fill.solid(); rect.fill.fore_color.rgb = RGBColor(0xFA, 0xFB, 0xFC)
            rect.line.color.rgb = RGBColor(0xB8, 0xC2, 0xCC); rect.line.dash_style = MSO_LINE_DASH_STYLE.DASH
            rect.shadow.inherit = False
            tf = rect.text_frame; tf.word_wrap = True
            first = (dg.get('text') or [''])[0]
            for j, (t, sz, b) in enumerate([(f"【Napkin 圖：{dg.get('en', '')}】", 16, True), (first, 14, False), ('（在 Napkin 生成後，刪除此框並貼上圖片）', 12, False)]):
                p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
                r = p.add_run(); r.text = t; r.font.size = Pt(sz); r.font.bold = b; r.font.color.rgb = RGBColor(0x5D, 0x68, 0x78)
                p.alignment = 2
        elif role == 'refs':
            c = by_type.get('refs', {}); k = 9
            refs = [f"{x['name']}：{x['url']}" for x in c.get('refs', [])]
            slots = sorted(bodies, key=lambda b: (_abs_pos(b)[1], _abs_pos(b)[0]))
            if slots:
                n = len(slots)
                for j, b in enumerate(slots):
                    bx = _abs_pos(b)[0]
                    if j < n - 1:
                        t = refs[j] if j < len(refs) else ''
                    else:
                        t = '\n'.join(refs[j:]) if j < len(refs) else ''
                    if b.width and (W - bx - margin) > b.width / EMU_IN:
                        b.width = Inches(W - bx - margin)
                    _set_runs(b, [[(line, {})] for line in t.split('\n')] if t else [[('', {})]], size=min(_font_size(b), 14))
                if len(refs) > n:
                    report.append(f'△ 參考資料有 {len(refs)} 筆，範本只有 {n} 個位置，多的放在最後一格')
        else:
            report.append(f"△ 第 {d['index'] + 1} 頁看不出用途（{d['title_text'][:10] or '沒有標題'}），沒有更動，不需要可以刪掉")
        if k is not None and k < len(script):
            s_obj.notes_slide.notes_text_frame.text = '\n'.join(f'{w}：{t}' for w, t in script[k].get('lines', []))
    prs.save(out_path)
    head = ['已依範本產生 ' + out_path, '頁面對應：' + '、'.join(
        f"第{d['index'] + 1}頁{ {'cover': '封面', 'toc': '目錄', 'intro': '前言', 'content': '內容', 'conclusion': '結語', 'refs': '參考', 'other': '（不動）'}[d['role']]}" for d in info)]
    return '\n'.join(head + report)


import sys as _sys
if __name__ == '__main__' and _sys.argv and _sys.argv[0].endswith('ppt_fill.py'):
    import json
    if len(_sys.argv) == 2:
        print(inspect_template(_sys.argv[1]))
    elif len(_sys.argv) >= 4:
        print(fill_template(_sys.argv[1], json.load(open(_sys.argv[2], encoding='utf-8')), _sys.argv[3]))
    else:
        print('用法：python3 ppt_fill.py 範本.pptx  或  python3 ppt_fill.py 範本.pptx deck.json 輸出.pptx')
